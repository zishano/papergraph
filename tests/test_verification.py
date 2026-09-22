from pathlib import Path

import httpx
import pytest
from openpyxl import load_workbook

from papergraph.models import CitationGraph
from papergraph.retrieval.openalex import OpenAlexClient
from papergraph.retrieval.verification import ArxivCandidateDiscovery, ArxivVerifier, keyword_terms
from papergraph.excel import write_excel


def graph():
    return CitationGraph.model_validate_json((Path(__file__).resolve().parents[1] /
        'validation/llmshare-dse/graph-citing.json').read_text())


def transport(seed_title, mode='match', oa_status=200):
    def handler(request):
        assert 'authorization' not in request.headers
        if request.url.host == 'api.openalex.org':
            return httpx.Response(oa_status, json={'id': 'https://openalex.org/W999', 'title': 'Candidate',
                 'doi': 'https://doi.org/10.48550/arxiv.2607.07096', 'referenced_works': []})
        if '/abs/' in request.url.path:
            return httpx.Response(200, text='<meta name="citation_title" content="Candidate"><meta name="citation_date" content="2026/07/08">')
        if mode == 'unavailable':
            return httpx.Response(404)
        body = '<h1 class="ltx_title">Candidate</h1>'
        if mode == 'body_only':
            body += f'<p>{seed_title}</p>'
        else:
            body += f'<li class="ltx_bibitem" id="bib.bib41">[41] {seed_title if mode == "match" else "Different paper"}</li>'
        return httpx.Response(200, text=body)
    return httpx.MockTransport(handler)


async def verify(g, mode='match', status=200):
    async with httpx.AsyncClient(transport=transport(g.nodes[0].title, mode, status)) as http:
        verifier = ArxivVerifier(http)
        records = await verifier.verify(g, [{'arxiv_id': '2607.07096', 'source': 'scholar_csv'},
                                           {'arxiv_id': '2607.07096v2'}], OpenAlexClient(http, api_key=''))
    return records


async def test_verified_reference_adds_direct_edge_and_excel_evidence(tmp_path):
    g = graph()
    records = await verify(g)
    assert len(records) == 1 and records[0]['status'] == 'verified'
    assert sum(n.hop == 1 for n in g.nodes) == 3
    assert any(e.source_id == 'W999' and e.target_id == g.seed_ids[0] and e.provider == 'arxiv_html' for e in g.edges)
    assert any(p.path == [g.seed_ids[0], 'W999'] for p in g.paths)
    book = load_workbook(write_excel(g, tmp_path/'test.xlsx', 'DSE', records))
    assert book['补充核验']['D2'].value == 'verified'
    assert '#bib.bib41' in book['补充核验']['E2'].value
    assert book['论文列表'].max_row == 6
    book.close()


@pytest.mark.parametrize('mode', ['body_only', 'unavailable', 'no_match'])
async def test_mentions_and_missing_html_never_create_edges(mode):
    g = graph()
    records = await verify(g, mode)
    assert len(g.nodes) == 5 and len(g.edges) == 5
    assert records[0]['status'] != 'verified'


async def test_arxiv_only_candidate_does_not_require_openalex():
    g = graph()
    records = await verify(g, status=404)
    assert records[0]['status'] == 'verified'
    assert next(n for n in g.nodes if n.id == 'arxiv:2607.07096').openalex_id is None


async def test_node_cap_retains_verified_evidence_without_adding_node():
    g = graph()
    g.config = g.config.model_copy(update={'max_papers': 5})
    records = await verify(g)
    assert records[0]['status'] == 'verified_not_added_node_limit'
    assert len(g.nodes) == 5
    assert g.truncated


async def test_arxiv_cannot_expand_scholar_membership():
    g = graph()
    seed = next(n for n in g.nodes if n.id in g.seed_ids)
    seed.id = 'scholar:seed'
    g.nodes = [seed]
    g.seed_ids = [seed.id]
    g.edges = []
    records = await verify(g)
    assert records[0]['status'] == 'verified_not_in_scholar_list'
    assert len(g.nodes) == 1 and not g.edges


async def test_keyword_discovery_expands_dse_and_returns_candidates():
    atom = '''<feed xmlns="http://www.w3.org/2005/Atom">
      <entry><id>https://arxiv.org/abs/2607.07096v1</id>
      <title>ThermoDSE: A Thermal-Aware Design Space Exploration</title></entry>
    </feed>'''
    def handler(request):
        assert 'DSE' in str(request.url) and 'Design' in str(request.url)
        return httpx.Response(200, text=atom)
    async with httpx.AsyncClient(transport=httpx.MockTransport(handler)) as http:
        candidates = await ArxivCandidateDiscovery(http).search('DSE', 5)
    assert keyword_terms('DSE') == ['DSE', 'Design Space Exploration']
    assert candidates == [{'arxiv_id': '2607.07096',
                           'title': 'ThermoDSE: A Thermal-Aware Design Space Exploration',
                           'source': 'arxiv_keyword_search'}]


async def test_keyword_discovery_falls_back_to_arxiv_search_page():
    page = '''<form></form><li class="arxiv-result">
      <p class="list-title"><a href="https://arxiv.org/abs/2607.07096v1">id</a></p>
      <p class="title">ThermoDSE: Design Space Exploration</p></li>'''
    def handler(request):
        if request.url.host == 'export.arxiv.org':
            return httpx.Response(503)
        assert request.url.params['searchtype'] == 'title'
        assert request.url.params['size'] == '25'
        return httpx.Response(200, text=page)
    async with httpx.AsyncClient(transport=httpx.MockTransport(handler)) as http:
        candidates = await ArxivCandidateDiscovery(http).search('DSE', 20)
    assert candidates[0]['arxiv_id'] == '2607.07096'
    assert candidates[0]['title'] == 'ThermoDSE: Design Space Exploration'


async def test_discovered_candidate_skips_abs_lookup_and_adds_edge():
    g = graph()
    def handler(request):
        if request.url.host == 'api.openalex.org':
            return httpx.Response(404)
        assert '/abs/' not in request.url.path
        return httpx.Response(200, text=(
            '<h1 class="ltx_title">ThermoDSE</h1>'
            f'<li class="ltx_bibitem">{g.nodes[0].title}</li>'))
    candidate = {'arxiv_id': '2607.07096', 'title': 'ThermoDSE',
                 'source': 'arxiv_keyword_search'}
    async with httpx.AsyncClient(transport=httpx.MockTransport(handler)) as http:
        records = await ArxivVerifier(http).verify(g, [candidate], OpenAlexClient(http, api_key=''))
    assert records[0]['status'] == 'verified'
    assert any(e.provider == 'arxiv_html' and e.target_id == g.seed_ids[0] for e in g.edges)
