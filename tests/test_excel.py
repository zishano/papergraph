from pathlib import Path

from openpyxl import load_workbook

from papergraph.cli import parser, run
from papergraph.excel import write_excel
from papergraph.models import CitationGraph


def sample():
    return CitationGraph.model_validate_json((Path(__file__).resolve().parents[1] /
        'validation/llmshare-dse/graph-citing.json').read_text())


def test_excel_preserves_nonmatching_papers_and_citation_direction(tmp_path):
    graph = sample()
    output = write_excel(graph, tmp_path / 'result.xlsx', 'DSE')
    book = load_workbook(output)
    ws = book['论文列表']
    assert ws.max_row == 5  # four discovered papers; seed has its own sheet
    assert [ws.cell(i, 3).value for i in range(2, 6)] == [
        '直接引用种子（Hop 1）', '直接引用种子（Hop 1）',
        '二跳及以上间接关联（Hop 2）', '二跳及以上间接关联（Hop 2）']
    assert all(ws.cell(i, 6).value == '未命中（不排除）' for i in range(2, 6))
    assert book['引用边'].max_row == 6
    assert book['发现路径'].max_row == 5
    assert book['add'].max_row == 5
    assert ws.freeze_panes == 'A2'
    assert ws.auto_filter.ref == 'A1:J5'
    assert book['引用边']['A2'].value == 'W4416429485'
    assert book['引用边']['C2'].value == 'W4414197293'
    book.close()


def test_metadata_is_literal_and_dse_expands(tmp_path):
    graph = sample()
    paper = next(n for n in graph.nodes if n.hop == 1)
    paper.title = '=1+1'
    paper.abstract = 'design-space exploration'
    book = load_workbook(write_excel(graph, tmp_path / 'literal.xlsx', 'DSE'))
    assert book['论文列表']['A2'].value == '=1+1'
    assert book['论文列表']['A2'].data_type == 's'
    assert book['论文列表']['F2'].value == 'DSE'
    book.close()


async def test_search_cli_writes_excel(tmp_path, monkeypatch):
    graph = sample()
    async def resolve(*args):
        return [n for n in graph.nodes if n.id in graph.seed_ids]
    async def crawl(*args):
        return graph
    monkeypatch.setattr('papergraph.cli.SeedResolver.resolve_many', resolve)
    monkeypatch.setattr('papergraph.cli.CitationCrawler.crawl', crawl)
    output = tmp_path / 'cli.xlsx'
    args = parser().parse_args(['search', '--source', 'openalex', '--seed', 'W4414197293',
                                '--keywords', 'DSE', '--no-arxiv-auto-verify',
                                '--excel', str(output)])
    result = await run(args)
    assert output.exists()
    assert len(result['nodes']) == 5
